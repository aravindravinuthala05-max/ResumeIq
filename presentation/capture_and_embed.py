import base64
import json
import os
import tempfile
import time
import urllib.request
from pathlib import Path

import websocket
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
PRESENTATION = ROOT / 'presentation' / 'ResumeIQ_Final_Presentation.pptx'
SHOT_DIR = ROOT / 'presentation' / 'screenshots'
SHOT_DIR.mkdir(exist_ok=True)

RESUME = '''Fictional Candidate
fictional.candidate@example.com | 555-010-1234
linkedin.com/in/fictional-candidate | github.com/fictional-candidate

SUMMARY
Backend developer focused on reliable web and data systems.

SKILLS
Python, JavaScript, TypeScript, Flask, React, SQL, PostgreSQL, AWS, Docker, Git

EXPERIENCE
Python Developer Intern at Example Labs
Software Engineer at Example Systems
Implemented APIs, improved response time by 30%, and optimized automation.

PROJECTS
Resume Search project: Flask, React, PostgreSQL, Docker, and machine learning.
Sensor project: Python, TensorFlow, and IoT with 100+ users.

EDUCATION
B.Tech in Computer Science, Fictional University, CGPA 8.5

CERTIFICATIONS
AWS Certified Developer, Docker Certification

ACHIEVEMENTS
Increased test coverage by 40%; reduced processing time by 2 seconds.

LANGUAGES
Python, JavaScript, English
'''
JD = 'Looking for a Python backend engineer with Flask, SQL, PostgreSQL, AWS, Docker, Git, and machine learning experience.'

def cdp(ws, method, params=None):
    cdp.n += 1
    ws.send(json.dumps({'id': cdp.n, 'method': method, 'params': params or {}}))
    while True:
        message = json.loads(ws.recv())
        if message.get('id') == cdp.n:
            if 'error' in message:
                raise RuntimeError(f"{method}: {message['error']}")
            return message.get('result', {})
cdp.n = 0

def evaluate(ws, expression):
    result = cdp(ws, 'Runtime.evaluate', {'expression': expression, 'returnByValue': True, 'awaitPromise': True})
    return result.get('result', {}).get('value')

def wait_for(ws, expression, timeout=20):
    end = time.time() + timeout
    while time.time() < end:
        if evaluate(ws, expression):
            return
        time.sleep(.25)
    raise TimeoutError(expression)

def screenshot(ws, filename, selector=None, max_height=None):
    params = {'format': 'png', 'captureBeyondViewport': True}
    if selector:
        rect = evaluate(ws, f'''(() => {{ const e=document.querySelector({json.dumps(selector)}); if(!e) return null; const r=e.getBoundingClientRect(); return {{x:r.left+scrollX-12,y:r.top+scrollY-12,width:r.width+24,height:r.height+24}}; }})()''')
        if not rect:
            raise RuntimeError(f'Missing selector: {selector}')
        if max_height:
            rect['height'] = min(rect['height'], max_height)
        rect['scale'] = 1
        params['clip'] = rect
    data = cdp(ws, 'Page.captureScreenshot', params)['data']
    (SHOT_DIR / filename).write_bytes(base64.b64decode(data))

def write_fixture():
    path = Path(tempfile.gettempdir()) / 'resumeiq-fictional-demo.pdf'
    doc = canvas.Canvas(str(path)); y = 800
    for line in RESUME.splitlines():
        doc.drawString(40, y, line[:110]); y -= 14
        if y < 40:
            doc.showPage(); y = 800
    doc.save()
    return path

def remove_placeholder(slide, label):
    for shape in list(slide.shapes):
        if getattr(shape, 'has_text_frame', False) and label in shape.text:
            element = shape._element
            element.getparent().remove(element)
            return True
    return False

def remove_picture_at(slide, x, y):
    for shape in list(slide.shapes):
        if getattr(shape, 'shape_type', None) == 13 and abs(shape.left - Inches(x)) < 1000 and abs(shape.top - Inches(y)) < 1000:
            element = shape._element
            element.getparent().remove(element)
            return True
    return False

def embed():
    prs = Presentation(PRESENTATION)
    # Remove all original placeholder labels and replace their visual spaces with actual UI screenshots.
    placements = [
        (4, 'Insert ResumeIQ Dashboard Screenshot', '01_landing.png', 7.4, 1.42, 4.9, 3.35),
        (10, 'Insert ATS / JD Match Screenshot', '02_dashboard.png', 8.35, 2.2, 3.3, 2.25),
        (11, 'Insert Skills Intelligence Screenshot', '03_skills.png', 6.75, 4.55, 5.2, 1.25),
        (13, 'Insert Interview Preparation Screenshot', '06_interview.png', 8.1, 5.48, 3.85, .72),
        (14, 'Insert ResumeIQ Dashboard Screenshot', '02_dashboard.png', .78, 1.35, 5.55, 3.95),
        (14, 'Insert PDF Report Screenshot', '01_landing.png', 6.95, 1.35, 5.55, 3.95),
    ]
    for index, label, image, x, y, w, h in placements:
        slide = prs.slides[index]
        if not remove_placeholder(slide, label):
            remove_picture_at(slide, x, y)
        slide.shapes.add_picture(str(SHOT_DIR / image), Inches(x), Inches(y), Inches(w), Inches(h))

    # Add the two remaining requested application views to their matching advisor/rewriter slide.
    slide = prs.slides[12]
    remove_picture_at(slide, .98, 5.12)
    remove_picture_at(slide, 6.98, 5.12)
    slide.shapes.add_picture(str(SHOT_DIR / '04_advisor.png'), Inches(.98), Inches(5.12), Inches(5.15), Inches(.64))
    slide.shapes.add_picture(str(SHOT_DIR / '05_rewriter.png'), Inches(6.98), Inches(5.12), Inches(5.15), Inches(.64))
    prs.save(PRESENTATION)

def main():
    endpoint = json.load(urllib.request.urlopen('http://127.0.0.1:9223/json'))[0]['webSocketDebuggerUrl']
    ws = websocket.create_connection(endpoint, timeout=30)
    cdp(ws, 'Page.enable'); cdp(ws, 'Runtime.enable')
    cdp(ws, 'Emulation.setDeviceMetricsOverride', {'width': 1440, 'height': 900, 'deviceScaleFactor': 1, 'mobile': False})
    cdp(ws, 'Page.navigate', {'url': 'http://127.0.0.1:5000/'})
    wait_for(ws, "Boolean(document.querySelector('#uploadForm'))")
    time.sleep(.8)
    screenshot(ws, '01_landing.png')
    fixture = write_fixture()
    doc = cdp(ws, 'DOM.getDocument', {'depth': 1})['root']['nodeId']
    input_id = cdp(ws, 'DOM.querySelector', {'nodeId': doc, 'selector': '#resume'})['nodeId']
    cdp(ws, 'DOM.setFileInputFiles', {'files': [str(fixture)], 'nodeId': input_id})
    evaluate(ws, f"(() => {{ const e=document.querySelector('#jobDescription'); e.value={json.dumps(JD)}; e.dispatchEvent(new Event('input',{{bubbles:true}})); document.querySelector('#uploadForm').requestSubmit(); return true; }})()")
    wait_for(ws, "location.pathname === '/result'", 30)
    wait_for(ws, "document.querySelector('#scoreValue') && document.querySelector('#scoreValue').textContent !== '0%'", 20)
    time.sleep(1)
    screenshot(ws, '02_dashboard.png', '.report-page', 1200)
    screenshot(ws, '03_skills.png', '#skills-section')
    screenshot(ws, '04_advisor.png', '#priority-improvements', 700)
    evaluate(ws, "showSection('rewriter'); true")
    wait_for(ws, "document.querySelector('#rewriter-section').dataset.loaded === 'true'", 20)
    screenshot(ws, '05_rewriter.png', '#rewriter-section', 720)
    screenshot(ws, '06_interview.png', '#interview-section', 850)
    ws.close()
    embed()

if __name__ == '__main__':
    main()
